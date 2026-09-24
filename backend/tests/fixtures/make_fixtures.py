"""Regenerates the binary test fixtures.

sample.docx/.pptx/.xlsx  built with python-docx / python-pptx / openpyxl (see build_office()).
sample.doc/.ppt/.xls     produced ONCE by converting the files above with LibreOffice
                         (`soffice --headless --convert-to doc|ppt|xls`) and committed. LibreOffice is only ever needed to
                         regenerate fixtures; neither the application nor the test suite uses it or any subprocess.
Everything else below needs only the normal dependencies.
"""
from __future__ import annotations

import json
from pathlib import Path

HERE = Path(__file__).parent


def build_office():
    import openpyxl
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    d = Document()
    d.add_heading("Gradient Descent Notes", 0)
    d.add_heading("Introduction", 1)
    d.add_paragraph("Gradient descent minimises a loss function by stepping against the gradient.")
    d.add_paragraph("Learning rate controls the step size.", style="List Bullet")
    d.add_heading("Table of methods", 2)
    t = d.add_table(rows=3, cols=2)
    for i, (a, b) in enumerate([("Method", "Note"), ("SGD", "uses one sample"), ("Adam", "adaptive moments")]):
        t.cell(i, 0).text, t.cell(i, 1).text = a, b
    d.add_paragraph("Résumé: naïve café — مرحبا بالعالم")
    d.save(HERE / "sample.docx")
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = "Backpropagation"
    s.placeholders[1].text = "Chain rule applies layer by layer\nCache activations"
    s.notes_slide.notes_text_frame.text = "Remember to derive the chain rule on the board"
    s2 = p.slides.add_slide(p.slide_layouts[5]); s2.shapes.title.text = "Loss surfaces"
    s2.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = "Free textbox: saddle points are common"
    p.save(HERE / "sample.pptx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Scores"; ws.append(["student", "score", "passed"])
    for i in range(1, 8):
        ws.append([f"stu{i}", 50 + i * 5, i > 3])
    wb.create_sheet("Notes")["A1"] = "Weights"
    wb.save(HERE / "sample.xlsx")


def build_simple():
    import pymupdf
    from PIL import Image, ImageDraw, ImageFont
    doc = pymupdf.open()
    for text in ("Gradient descent updates parameters using the learning rate.\nThe update rule is w = w - lr * gradient.",
                 "Momentum accumulates past gradients to smooth the optimisation path."):
        page = doc.new_page(); page.insert_text((72, 90), text, fontsize=12)
    doc.save(HERE / "sample.pdf")
    img = Image.new("RGB", (1000, 240), "white"); d = ImageDraw.Draw(img)
    f = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 36)
    d.text((30, 30), "Backpropagation uses the chain rule", fill="black", font=f)
    d.text((30, 130), "Assignment due Friday", fill="black", font=f)
    img.save(HERE / "sample.png")
    (HERE / "sample.csv").write_text("student,score\nAda,91\nMohab,88\nÉlodie,95\n", encoding="utf-8")
    (HERE / "sample.txt").write_text("Regularisation reduces overfitting by penalising large weights.\n", encoding="utf-8")
    (HERE / "sample.md").write_text("# Overfitting\n\nA model that memorises noise generalises badly.\n\n## Remedies\n- more data\n- regularisation\n", encoding="utf-8")
    (HERE / "sample.py").write_text('"""Linear regression demo."""\n\ndef predict(w, x):\n    """Return w * x."""\n    return w * x\n', encoding="utf-8")
    (HERE / "sample.ipynb").write_text(json.dumps({"cells": [{"cell_type": "markdown", "source": ["# Lab 1"]}, {"cell_type": "code", "source": ["print(1+1)"], "outputs": [{"output_type": "stream", "name": "stdout", "text": ["2\n"]}]}],
                                                    "metadata": {"language_info": {"name": "python"}}, "nbformat": 4, "nbformat_minor": 5}), encoding="utf-8")


if __name__ == "__main__":
    build_simple()
    build_office()
