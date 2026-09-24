"""Format-specific extraction -> normalised Markdown pages.

Every extractor takes validated bytes and returns Extracted(pages=[(page_no, markdown)]). Headings become `#`, code
becomes fenced blocks, tables become Markdown tables: that is the structure the RAG chunker already understands, so
nothing downstream needs to know the source format. Extractors never execute content and enforce size limits.
"""
from __future__ import annotations

import ast
import csv
import datetime as dt
import io
import json
from dataclasses import dataclass, field
from typing import Callable

from . import registry as R
from .errors import ExtractionError
from .textutil import check_size, decode_text, md_table, normalize_text

MAX_CELLS = 500_000
MAX_CSV_ROWS = 200_000
MAX_SLIDES = 500
MAX_NB_OUTPUT = 2_000


@dataclass
class Extracted:
    pages: list[tuple[int, str]]
    title: str | None = None
    metadata: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(t for _, t in self.pages)


def _finish(pages: list[tuple[int, str]], **kw) -> Extracted:
    pages = [(n, normalize_text(t)) for n, t in pages]
    pages = [(n, t) for n, t in pages if t]
    check_size("\n".join(t for _, t in pages))
    return Extracted(pages=pages, **kw)


# ------------------------------------------------------------------ PDF (plain text; RAG ingestion of PDFs uses the AI's layout-aware parser)
def _pdf(data: bytes, spec) -> Extracted:
    import pymupdf
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"PDF cannot be read: {exc}") from exc
    try:
        pages = [(i + 1, p.get_text("text")) for i, p in enumerate(doc)]
        title = (doc.metadata or {}).get("title") or None
    finally:
        doc.close()
    ex = _finish(pages, title=title, metadata={"pages": len(pages)})
    if not ex.pages:
        ex.warnings.append("no_text_layer")            # scanned PDF: needs OCR (AI's Docling fallback handles this for RAG)
    return ex


# ------------------------------------------------------------------ Word
def _docx(data: bytes, spec) -> Extracted:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Word document cannot be read: {exc}") from exc
    out: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style == "Title":
                out.append(f"# {text}")
            elif style.startswith("Heading "):
                try:
                    out.append("#" * min(6, max(1, int(style.split()[-1]))) + f" {text}")
                except ValueError:
                    out.append(text)
            elif style.startswith("List"):
                out.append(f"- {text}")
            else:
                out.append(text)
        elif tag == "tbl":
            t = Table(child, doc)
            out += md_table([[c.text for c in r.cells] for r in t.rows])
    cp = doc.core_properties
    return _finish([(1, "\n\n".join(out))], title=(cp.title or None), metadata={"author_present": bool(cp.author)})


def _doc(data: bytes, spec) -> Extracted:
    from .legacy_office import read_doc
    text = read_doc(data)
    return _finish([(1, text)], metadata={"legacy": True})


# ------------------------------------------------------------------ PowerPoint
def _walk_shapes(shapes):
    for sh in sorted(shapes, key=lambda s: (getattr(s, "top", 0) or 0, getattr(s, "left", 0) or 0)):
        if sh.shape_type == 6:                              # MSO_SHAPE_TYPE.GROUP
            yield from _walk_shapes(sh.shapes)
        else:
            yield sh


def _pptx(data: bytes, spec) -> Extracted:
    from pptx import Presentation
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"PowerPoint file cannot be read: {exc}") from exc
    slides = list(prs.slides)
    if len(slides) > MAX_SLIDES:
        raise ExtractionError(f"presentation has {len(slides)} slides (limit {MAX_SLIDES})", code="too_large")
    pages = []
    for i, slide in enumerate(slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        parts = [f"## Slide {i}: {title}" if title else f"## Slide {i}"]
        for sh in _walk_shapes(slide.shapes):
            if slide.shapes.title is not None and sh.shape_id == slide.shapes.title.shape_id:
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                parts += md_table([[c.text for c in r.cells] for r in sh.table.rows])
            elif sh.has_text_frame and sh.text_frame.text.strip():
                lines = [pp.text.strip() for pp in sh.text_frame.paragraphs if pp.text.strip()]
                parts.append("\n".join(f"- {l}" if len(lines) > 1 else l for l in lines))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        pages.append((i, "\n\n".join(parts)))
    return _finish(pages, metadata={"slides": len(slides)})


def _ppt(data: bytes, spec) -> Extracted:
    from .legacy_office import read_ppt
    pages = [(n, f"## Slide {n}\n\n{t}") for n, t in read_ppt(data)]
    return _finish(pages, metadata={"legacy": True}, warnings=["legacy_ppt_tables_and_notes_are_flattened_or_omitted"])


# ------------------------------------------------------------------ spreadsheets
def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return f"{v:.10g}"
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    return str(v)


def _sheet_page(idx: int, name: str, rows: list[list]) -> tuple[int, str]:
    blocks = md_table([[_fmt(c) for c in r] for r in rows])
    if not blocks:
        return idx, ""
    head = f"## Sheet: {name}"
    return idx, "\n\n".join([head] + [f"{b}" for b in blocks])


def _xlsx(data: bytes, spec) -> Extracted:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Excel workbook cannot be read: {exc}") from exc
    pages, cells, skipped = [], 0, []
    try:
        for idx, ws in enumerate(wb.worksheets, start=1):
            if ws.sheet_state != "visible":
                skipped.append(ws.title)
                continue
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells += len(row)
                if cells > MAX_CELLS:
                    raise ExtractionError(f"workbook has more than {MAX_CELLS:,} cells; split it", code="too_large")
                rows.append(list(row))
            pages.append(_sheet_page(idx, ws.title, rows))
    finally:
        wb.close()
    return _finish(pages, metadata={"sheets": len(pages), "hidden_sheets_skipped": skipped})


def _xls(data: bytes, spec) -> Extracted:
    import xlrd
    try:
        wb = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Excel workbook cannot be read: {exc}") from exc
    pages, cells, skipped = [], 0, []
    try:
        for idx in range(wb.nsheets):
            sh = wb.sheet_by_index(idx)
            if sh.visibility != 0:
                skipped.append(sh.name)
                continue
            rows = []
            for r in range(sh.nrows):
                cells += sh.ncols
                if cells > MAX_CELLS:
                    raise ExtractionError(f"workbook has more than {MAX_CELLS:,} cells; split it", code="too_large")
                vals = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            vals.append(xlrd.xldate.xldate_as_datetime(cell.value, wb.datemode))
                            continue
                        except Exception:  # noqa: BLE001
                            pass
                    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        vals.append(bool(cell.value))
                    elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                        vals.append(int(cell.value))
                    else:
                        vals.append(cell.value)
                rows.append(vals)
            pages.append(_sheet_page(idx + 1, sh.name, rows))
            wb.unload_sheet(idx)
    finally:
        wb.release_resources()
    return _finish(pages, metadata={"sheets": len(pages), "hidden_sheets_skipped": skipped})


def _csv(data: bytes, spec) -> Extracted:
    text, enc = decode_text(data)
    delimiter = "\t" if spec.ext == "tsv" else None
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:20_000], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    csv.field_size_limit(1_000_000)
    rows = []
    try:
        for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter)):
            if i >= MAX_CSV_ROWS:
                raise ExtractionError(f"table has more than {MAX_CSV_ROWS:,} rows; split it", code="too_large")
            rows.append(row)
    except csv.Error as exc:
        raise ExtractionError(f"CSV cannot be parsed: {exc}") from exc
    blocks = md_table(rows)
    return _finish([(1, "\n\n".join(blocks))], metadata={"rows": max(0, len(rows) - 1), "delimiter": delimiter, "encoding": enc})


# ------------------------------------------------------------------ text / markdown / code
def _text(data: bytes, spec) -> Extracted:
    text, enc = decode_text(data)
    return _finish([(1, text)], metadata={"encoding": enc})


def _python_outline(src: str) -> str:
    """Safe, static outline: ast.parse never executes code. Deep nesting or huge input just yields no outline."""
    if len(src) > 1_000_000:
        return ""
    try:
        tree = ast.parse(src)
    except (SyntaxError, ValueError, RecursionError, MemoryError):
        return ""
    items = []
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            kind = "class" if isinstance(node, ast.ClassDef) else "def"
            doc = (ast.get_docstring(node) or "").strip().split("\n")[0][:120]
            items.append(f"- {kind} {node.name}" + (f": {doc}" if doc else ""))
    return "Outline:\n" + "\n".join(items) if items else ""


def _code(data: bytes, spec) -> Extracted:
    text, enc = decode_text(data)
    lang = spec.language or ""
    parts = []
    if spec.ext == "py":
        outline = _python_outline(text)
        if outline:
            parts.append(outline)
    fence = "````" if "```" in text else "```"
    parts.append(f"{fence}{lang}\n{text.strip()}\n{fence}")
    return _finish([(1, "\n\n".join(parts))], metadata={"encoding": enc, "language": lang})


def _ipynb(data: bytes, spec) -> Extracted:
    text, _ = decode_text(data)
    try:
        nb = json.loads(text)
    except ValueError as exc:
        raise ExtractionError(f"notebook is not valid JSON: {exc}") from exc
    lang = ((nb.get("metadata") or {}).get("language_info") or {}).get("name") or "python"
    parts = []
    for cell in nb.get("cells", []):
        src = cell.get("source", "")
        src = "".join(src) if isinstance(src, list) else str(src)
        if not src.strip():
            continue
        if cell.get("cell_type") == "markdown":
            parts.append(src.strip())
        elif cell.get("cell_type") == "code":
            parts.append(f"```{lang}\n{src.strip()}\n```")
            for o in cell.get("outputs", []) or []:
                t = o.get("text") or (o.get("data") or {}).get("text/plain")      # never images/html/base64
                if t:
                    t = "".join(t) if isinstance(t, list) else str(t)
                    parts.append("Output:\n```\n" + t.strip()[:MAX_NB_OUTPUT] + "\n```")
    return _finish([(1, "\n\n".join(parts))], metadata={"cells": len(nb.get("cells", [])), "language": lang})


# ------------------------------------------------------------------ images
def _image(data: bytes, spec) -> Extracted:
    from .ocr import ocr_image
    text = ocr_image(data)
    ex = _finish([(1, text)] if text.strip() else [], metadata={"ocr": True})
    if not ex.pages:
        ex.warnings.append("no_text_detected")
    return ex


EXTRACTORS: dict[str, Callable] = {
    "pdf": _pdf, "docx": _docx, "doc": _doc, "pptx": _pptx, "ppt": _ppt, "xlsx": _xlsx, "xls": _xls, "csv": _csv,
    "text": _text, "markdown": _text, "code": _code, "ipynb": _ipynb, "image_ocr": _image,
}


def extract(spec: R.FormatSpec, data: bytes) -> Extracted:
    if spec.extractor is None:
        raise ExtractionError("no extractor for this format", code="unsupported")
    return EXTRACTORS[spec.extractor](data, spec)
